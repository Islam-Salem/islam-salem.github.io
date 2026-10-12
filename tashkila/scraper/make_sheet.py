#!/usr/bin/env python3
"""Writes the review sheet (review.xlsx, to upload to Google Drive as a Google Sheet) and
data/xi/sheet_base.json (what every cell said at the start, so only your changes are applied).

Usage: python scraper/make_sheet.py [out.xlsx]
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sheet as S  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "review.xlsx"
    d = os.path.join(ROOT, "data", "xi")
    with open(os.path.join(d, "data.json"), encoding="utf-8") as f:
        data = json.load(f)
    with open(os.path.join(d, "full.json"), encoding="utf-8") as f:
        data["criteria"] = json.load(f)
    rows = S.rows_from_data(data)

    wb = Workbook()
    ws = wb.active
    ws.title = "اللاعبون"
    ws.sheet_view.rightToLeft = True
    ws.append(S.HEADERS)
    for r in rows:
        ws.append([r.get(h, "") for h in S.HEADERS])
    head = PatternFill("solid", fgColor="0F3220")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"
    widths = {S.COL_ID: 13, S.COL_NAME: 24, S.COL_EN: 24, S.COL_IN: 9, S.COL_COUNTRIES: 34, S.COL_FCLUBS: 34,
              S.COL_NOTES: 30}
    grey = PatternFill("solid", fgColor="EEEEEE")
    tick_cols = [c for c, _ in S.POS_COLS + S.CLUB_COLS] + [c for c, _, _ in S.HON_COLS]
    n = len(rows) + 400   # room for players you add
    for k, h in enumerate(S.HEADERS, start=1):
        letter = ws.cell(1, k).column_letter
        ws.column_dimensions[letter].width = widths.get(h, 10)
        if h in tick_cols:
            dv = DataValidation(type="list", formula1=f'"{S.YES}"', allow_blank=True)
            dv.add(f"{letter}2:{letter}{n}")
            ws.add_data_validation(dv)
            for row in range(2, len(rows) + 2):
                ws.cell(row, k).alignment = Alignment(horizontal="center")
        if h == S.COL_IN:
            dv = DataValidation(type="list", formula1='"نعم,لا"', allow_blank=True)
            dv.add(f"{letter}2:{letter}{n}")
            ws.add_data_validation(dv)
        if h in (S.COL_ID, S.COL_EN):
            for row in range(2, len(rows) + 2):
                ws.cell(row, k).fill = grey

    help_ws = wb.create_sheet("طريقة الاستخدام")
    help_ws.sheet_view.rightToLeft = True
    lines = [
        "مراجعة لاعبي «تشكيلة المنتخب»",
        "",
        "كل صف لاعب. عدّل ما تريد في الصفحة الأولى «اللاعبون»، وتعديلاتك تدخل اللعبة مع التحديث الأسبوعي (أو عند تشغيل Update data).",
        "",
        "• لرفض لاعب: اكتب «لا» في عمود «في اللعبة».",
        "• لتعديل المركز أو الأهلي أو الزمالك أو البطولات: ضع ✓ أو امسحها.",
        "• لتعديل البلاد أو الأندية خارج مصر: اكتبها مفصولة بفاصلة، مثل: السعودية، الإمارات",
        "  (اكتب اسم البلد كما في صفحة «البلاد والأندية» حتى يظهر علمه).",
        "• لتغيير الاسم الذي يظهر في اللعبة: عدّل عمود «الاسم» (الاسم القديم يظل يعمل في البحث).",
        "• لإضافة لاعب: اكتبه في صف جديد في آخر الصفحة. رقم ويكي بيانات اختياري.",
        "",
        "الخانة التي لا تلمسها تتبع البيانات المحدثة من ويكيبيديا، والخانة التي عدلتها تبقى كما كتبتها.",
        "لا تغيّر ترتيب الصفحات ولا أسماء الأعمدة، ولا تمسح عمود «رقم ويكي بيانات».",
    ]
    for line in lines:
        help_ws.append([line])
    help_ws.column_dimensions["A"].width = 120
    help_ws["A1"].font = Font(bold=True, size=14)

    ref = wb.create_sheet("البلاد والأندية")
    ref.sheet_view.rightToLeft = True
    ref.append(["البلاد", "أندية خارج مصر", "البطولات"])
    countries = sorted({S.country_name(c["l"]) for c in data["criteria"].values() if c["t"] == "abroad"})
    clubs = sorted({c["l"] for c in data["criteria"].values() if c["t"] == "fclub"})
    hons = [label for _, _, label in S.HON_COLS]
    for k in range(max(len(countries), len(clubs), len(hons))):
        ref.append([countries[k] if k < len(countries) else "", clubs[k] if k < len(clubs) else "",
                    hons[k] if k < len(hons) else ""])
    for c in ref[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
    for col in "ABC":
        ref.column_dimensions[col].width = 30

    wb.save(out)
    with open(os.path.join(d, "sheet_base.json"), "w", encoding="utf-8") as f:
        json.dump(S.base_of(rows), f, ensure_ascii=False, separators=(",", ":"))
    print(f"{out}: {len(rows)} players")


if __name__ == "__main__":
    main()
